"""
Naija Pocket Business Center
Payment + Saved Document + Customer Care Delivery API

Complete replacement.

Core rules:
- Product identity = service + document title.
- IDs are internal only.
- The saved document is never regenerated during download.
- Download stays locked until Customer Care activates it.
- Back Office key = NPBC-2026.
- Delivery channels:
    1. Download to Phone
    2. WhatsApp
    3. Email
    4. Telegram
    5. Google Drive
- WhatsApp, Email and Telegram share the exact saved-document download link.
- Google Drive opens Drive for manual upload of the exact downloaded document.
"""

from __future__ import annotations

import os
import re
import sqlite3
import urllib.parse
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import FastAPI, Header, HTTPException, Query, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel


# ============================================================
# OPTIONAL DOCX SUPPORT
# ============================================================

try:
    from docx import Document
except Exception:
    Document = None


# ============================================================
# CONFIGURATION
# ============================================================

APP_VERSION = "payment-product-first-v5-delivery"

BASE_DIR = Path(__file__).resolve().parent

DOWNLOAD_DIR = BASE_DIR / "downloads"
DOWNLOAD_DIR.mkdir(parents=True, exist_ok=True)

PRODUCT_DB_PATH = BASE_DIR / "product_delivery.db"
PAYMENT_DB_PATH = BASE_DIR / "payment_gateway.db"

# SIMPLE FIXED BACK OFFICE KEY
BACK_OFFICE_ADMIN_KEY = "NPBC-2026"

# Optional legacy webhook support.
# These are NOT required for the normal customer delivery buttons.
DELIVERY_EMAIL_WEBHOOK = os.getenv(
    "DELIVERY_EMAIL_WEBHOOK", ""
).strip()

DELIVERY_WHATSAPP_WEBHOOK = os.getenv(
    "DELIVERY_WHATSAPP_WEBHOOK", ""
).strip()

DELIVERY_TELEGRAM_WEBHOOK = os.getenv(
    "DELIVERY_TELEGRAM_WEBHOOK", ""
).strip()

DELIVERY_GOOGLE_DRIVE_WEBHOOK = os.getenv(
    "DELIVERY_GOOGLE_DRIVE_WEBHOOK", ""
).strip()

# Optional public API URL.
#
# If this is not configured, the API builds the download URL from
# the current request host.
PUBLIC_API_BASE_URL = os.getenv(
    "PUBLIC_API_BASE_URL", ""
).strip().rstrip("/")

CURRENCY = "NGN"


# ============================================================
# FASTAPI
# ============================================================

app = FastAPI(
    title="Naija Pocket Business Center Payment API",
    version=APP_VERSION,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ============================================================
# GENERAL HELPERS
# ============================================================

def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def clean_text(value: Any) -> str:
    if value is None:
        return ""

    return str(value).strip()


def normalize_key_part(value: Any) -> str:
    value = clean_text(value).casefold()
    value = re.sub(r"\s+", " ", value)
    return value


def business_key(service: Any, title: Any) -> str:
    return (
        f"{normalize_key_part(service)}::"
        f"{normalize_key_part(title)}"
    )


def safe_filename(
    value: Any,
    fallback: str = "document",
) -> str:

    name = clean_text(value)

    name = Path(name).name

    if not name:
        name = fallback

    name = re.sub(
        r"[^\w.\- ()]+",
        "_",
        name,
        flags=re.UNICODE,
    )

    name = re.sub(
        r"\s+",
        " ",
        name,
    ).strip()

    return name[:180] or fallback


def safe_folder_name(
    service: Any,
    title: Any,
) -> str:

    raw = (
        f"{clean_text(service)} - "
        f"{clean_text(title)}"
    )

    raw = re.sub(
        r"[^\w\- ]+",
        "_",
        raw,
        flags=re.UNICODE,
    )

    raw = re.sub(
        r"\s+",
        " ",
        raw,
    ).strip()

    return raw[:180] or "document"


def first_nonempty(
    data: Dict[str, Any],
    *keys: str,
    default: Any = None,
) -> Any:

    for key in keys:
        if key in data:
            value = data[key]

            if value not in (
                None,
                "",
                [],
                {},
            ):
                return value

    return default


def db_connect(path: Path) -> sqlite3.Connection:
    conn = sqlite3.connect(str(path))
    conn.row_factory = sqlite3.Row
    return conn


def row_dict(
    row: Optional[sqlite3.Row],
) -> Optional[Dict[str, Any]]:

    if row is None:
        return None

    return {
        key: row[key]
        for key in row.keys()
    }


# ============================================================
# DATABASE INITIALIZATION
# ============================================================

def init_product_db() -> None:

    with db_connect(PRODUCT_DB_PATH) as conn:

        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS document_products (
                id INTEGER PRIMARY KEY AUTOINCREMENT,

                business_key TEXT NOT NULL UNIQUE,

                service TEXT NOT NULL,

                document_title TEXT NOT NULL,

                customer_name TEXT,

                customer_id TEXT,

                amount REAL DEFAULT 0,

                currency TEXT DEFAULT 'NGN',

                document_version TEXT,

                document_filename TEXT,

                document_pages INTEGER DEFAULT 1,

                document_payload TEXT,

                document_saved_path TEXT,

                document_saved_at TEXT,

                download_unlocked INTEGER DEFAULT 0,

                created_at TEXT NOT NULL,

                updated_at TEXT NOT NULL,

                activated_at TEXT,

                downloaded_at TEXT,

                download_count INTEGER DEFAULT 0,

                notes TEXT
            )
            """
        )

        # Delivery history is separate so existing product/payment
        # records do not need to be rebuilt.
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS delivery_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,

                business_key TEXT NOT NULL,

                service TEXT NOT NULL,

                document_title TEXT NOT NULL,

                channel TEXT NOT NULL,

                recipient TEXT,

                action_url TEXT,

                status TEXT NOT NULL,

                message TEXT,

                created_at TEXT NOT NULL
            )
            """
        )

        conn.commit()


def init_payment_db() -> None:

    with db_connect(PAYMENT_DB_PATH) as conn:

        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS payment_orders (
                id INTEGER PRIMARY KEY AUTOINCREMENT,

                business_key TEXT NOT NULL UNIQUE,

                service TEXT NOT NULL,

                document_title TEXT NOT NULL,

                customer_name TEXT,

                customer_id TEXT,

                amount REAL DEFAULT 0,

                currency TEXT DEFAULT 'NGN',

                payment_method TEXT,

                payment_status TEXT DEFAULT 'created',

                document_version TEXT,

                document_filename TEXT,

                document_pages INTEGER DEFAULT 1,

                notes TEXT,

                created_at TEXT NOT NULL,

                updated_at TEXT NOT NULL,

                reported_at TEXT,

                verified_at TEXT
            )
            """
        )

        conn.commit()


def init_databases() -> None:
    init_product_db()
    init_payment_db()


# ============================================================
# DOCUMENT NORMALIZATION
# ============================================================

def normalize_pages(value: Any) -> List[str]:

    if value is None:
        return []

    if isinstance(value, str):

        text = value.strip()

        if not text:
            return []

        return [text]

    if isinstance(value, list):

        pages = []

        for item in value:

            if isinstance(item, dict):

                item_text = first_nonempty(
                    item,
                    "text",
                    "content",
                    "page_text",
                    default="",
                )

            else:

                item_text = item

            item_text = clean_text(item_text)

            if item_text:
                pages.append(item_text)

        return pages

    if isinstance(value, dict):

        text = first_nonempty(
            value,
            "text",
            "content",
            "document_text",
            default="",
        )

        text = clean_text(text)

        return [text] if text else []

    text = clean_text(value)

    return [text] if text else []


def normalize_document_payload(
    data: Dict[str, Any],
) -> Dict[str, Any]:

    pages = normalize_pages(
        first_nonempty(
            data,
            "pages",
            "document_pages",
            "page_text",
            default=None,
        )
    )

    document_text = clean_text(
        first_nonempty(
            data,
            "document_text",
            "documentText",
            "text",
            "content",
            default="",
        )
    )

    if not document_text and pages:
        document_text = "\n\n".join(pages)

    if not pages and document_text:
        pages = [document_text]

    filename = safe_filename(
        first_nonempty(
            data,
            "filename",
            "document_filename",
            "documentFilename",
            default="Naija_Pocket_Document.docx",
        ),
        fallback="Naija_Pocket_Document.docx",
    )

    if not filename.lower().endswith(
        (
            ".docx",
            ".pdf",
        )
    ):
        filename += ".docx"

    version = clean_text(
        first_nonempty(
            data,
            "document_version",
            "version_id",
            "versionId",
            "documentVersion",
            default="",
        )
    )

    return {
        "text": document_text,
        "pages": pages,
        "filename": filename,
        "version": version,
        "page_count": len(pages) or 1,
    }


# ============================================================
# DOCX CREATION
# ============================================================

def make_docx(
    text: str,
    pages: List[str],
    output_path: Path,
) -> None:

    if Document is None:

        raise RuntimeError(
            "python-docx is not installed. "
            "Install it with: pip install python-docx"
        )

    document = Document()

    actual_pages = pages or [
        text
    ]

    for index, page_text in enumerate(
        actual_pages
    ):

        if index:
            document.add_page_break()

        lines = str(page_text).splitlines()

        if not lines:
            document.add_paragraph("")
            continue

        for line in lines:
            document.add_paragraph(line)

    document.save(
        str(output_path)
    )


# ============================================================
# SAVE EXACT DOCUMENT SNAPSHOT
# ============================================================

def save_exact_document_snapshot(
    service: str,
    title: str,
    payload: Dict[str, Any],
) -> str:

    key = business_key(
        service,
        title,
    )

    # Important:
    # NEVER replace an already saved document.
    with db_connect(PRODUCT_DB_PATH) as conn:

        existing = conn.execute(
            """
            SELECT document_saved_path
            FROM document_products
            WHERE business_key = ?
            """,
            (key,),
        ).fetchone()

        if existing:

            existing_path = clean_text(
                existing["document_saved_path"]
            )

            if existing_path:

                existing_file = Path(
                    existing_path
                )

                if (
                    existing_file.exists()
                    and existing_file.is_file()
                ):
                    return str(
                        existing_file
                    )

    folder = (
        DOWNLOAD_DIR
        / safe_folder_name(
            service,
            title,
        )
    )

    folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    filename = safe_filename(
        payload["filename"]
    )

    output_path = (
        folder
        / filename
    )

    # If the same exact snapshot already exists,
    # keep it.
    if (
        output_path.exists()
        and output_path.is_file()
    ):

        saved_path = str(
            output_path
        )

    else:

        make_docx(
            payload["text"],
            payload["pages"],
            output_path,
        )

        saved_path = str(
            output_path
        )

    with db_connect(PRODUCT_DB_PATH) as conn:

        conn.execute(
            """
            UPDATE document_products

            SET document_saved_path = ?,

                document_saved_at =
                    COALESCE(
                        document_saved_at,
                        ?
                    ),

                updated_at = ?

            WHERE business_key = ?
            """,
            (
                saved_path,
                now_iso(),
                now_iso(),
                key,
            ),
        )

        conn.commit()

    return saved_path


# ============================================================
# PRODUCT LOOKUPS
# ============================================================

def get_product(
    service: str,
    title: str,
) -> Optional[Dict[str, Any]]:

    key = business_key(
        service,
        title,
    )

    with db_connect(PRODUCT_DB_PATH) as conn:

        row = conn.execute(
            """
            SELECT *
            FROM document_products
            WHERE business_key = ?
            """,
            (key,),
        ).fetchone()

    return row_dict(row)


def get_single_product() -> Optional[Dict[str, Any]]:

    with db_connect(PRODUCT_DB_PATH) as conn:

        row = conn.execute(
            """
            SELECT *
            FROM document_products

            ORDER BY
                updated_at DESC,
                id DESC

            LIMIT 1
            """
        ).fetchone()

    return row_dict(row)


def get_product_or_single(
    service: str,
    title: str,
) -> Optional[Dict[str, Any]]:

    product = get_product(
        service,
        title,
    )

    if product:
        return product

    # Compatibility fallback for older
    # single-document records.
    return get_single_product()


# ============================================================
# UPSERT PRODUCT
# ============================================================

def upsert_product(
    service: str,
    title: str,
    customer_name: str,
    customer_id: str,
    amount: float,
    currency: str,
    document: Dict[str, Any],
    notes: str = "",
) -> Dict[str, Any]:

    key = business_key(
        service,
        title,
    )

    current_time = now_iso()

    with db_connect(PRODUCT_DB_PATH) as conn:

        existing = conn.execute(
            """
            SELECT *
            FROM document_products
            WHERE business_key = ?
            """,
            (key,),
        ).fetchone()

        if existing:

            conn.execute(
                """
                UPDATE document_products

                SET customer_name = ?,

                    customer_id = ?,

                    amount = ?,

                    currency = ?,

                    document_version = ?,

                    document_filename = ?,

                    document_pages = ?,

                    document_payload =
                        CASE
                            WHEN
                                document_saved_path IS NULL
                                OR document_saved_path = ''
                            THEN ?
                            ELSE document_payload
                        END,

                    updated_at = ?,

                    notes = ?

                WHERE business_key = ?
                """,
                (
                    customer_name,
                    customer_id,
                    amount,
                    currency,
                    document["version"],
                    document["filename"],
                    document["page_count"],
                    document["text"],
                    current_time,
                    notes,
                    key,
                ),
            )

        else:

            conn.execute(
                """
                INSERT INTO document_products (

                    business_key,

                    service,

                    document_title,

                    customer_name,

                    customer_id,

                    amount,

                    currency,

                    document_version,

                    document_filename,

                    document_pages,

                    document_payload,

                    created_at,

                    updated_at,

                    notes

                )

                VALUES (
                    ?, ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?, ?, ?
                )
                """,
                (
                    key,
                    service,
                    title,
                    customer_name,
                    customer_id,
                    amount,
                    currency,
                    document["version"],
                    document["filename"],
                    document["page_count"],
                    document["text"],
                    current_time,
                    current_time,
                    notes,
                ),
            )

        conn.commit()

        row = conn.execute(
            """
            SELECT *
            FROM document_products
            WHERE business_key = ?
            """,
            (key,),
        ).fetchone()

    return row_dict(row) or {}


# ============================================================
# ACTIVATE DOWNLOAD
# ============================================================

def activate_product_download(
    service: str,
    title: str,
) -> Dict[str, Any]:

    product = get_product_or_single(
        service,
        title,
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    key = business_key(
        product["service"],
        product["document_title"],
    )

    activated_at = now_iso()

    with db_connect(PRODUCT_DB_PATH) as conn:

        conn.execute(
            """
            UPDATE document_products

            SET download_unlocked = 1,

                activated_at = ?,

                updated_at = ?

            WHERE business_key = ?
            """,
            (
                activated_at,
                activated_at,
                key,
            ),
        )

        conn.commit()

    return (
        get_product(
            product["service"],
            product["document_title"],
        )
        or product
    )


# ============================================================
# PAYMENT LOOKUPS
# ============================================================

def get_payment(
    service: str,
    title: str,
) -> Optional[Dict[str, Any]]:

    key = business_key(
        service,
        title,
    )

    with db_connect(PAYMENT_DB_PATH) as conn:

        row = conn.execute(
            """
            SELECT *
            FROM payment_orders
            WHERE business_key = ?
            """,
            (key,),
        ).fetchone()

    return row_dict(row)


# ============================================================
# PAYMENT RECORD
# ============================================================

def ensure_payment_record(
    service: str,
    title: str,
    customer_name: str,
    customer_id: str,
    amount: float,
    currency: str,
    payment_method: str,
    document: Dict[str, Any],
) -> Dict[str, Any]:

    key = business_key(
        service,
        title,
    )

    current_time = now_iso()

    with db_connect(PAYMENT_DB_PATH) as conn:

        existing = conn.execute(
            """
            SELECT *
            FROM payment_orders
            WHERE business_key = ?
            """,
            (key,),
        ).fetchone()

        if existing:

            conn.execute(
                """
                UPDATE payment_orders

                SET customer_name = ?,

                    customer_id = ?,

                    amount = ?,

                    currency = ?,

                    payment_method = ?,

                    document_version = ?,

                    document_filename = ?,

                    document_pages = ?,

                    updated_at = ?

                WHERE business_key = ?
                """,
                (
                    customer_name,
                    customer_id,
                    amount,
                    currency,
                    payment_method,
                    document["version"],
                    document["filename"],
                    document["page_count"],
                    current_time,
                    key,
                ),
            )

        else:

            conn.execute(
                """
                INSERT INTO payment_orders (

                    business_key,

                    service,

                    document_title,

                    customer_name,

                    customer_id,

                    amount,

                    currency,

                    payment_method,

                    payment_status,

                    document_version,

                    document_filename,

                    document_pages,

                    created_at,

                    updated_at

                )

                VALUES (
                    ?, ?, ?, ?, ?, ?, ?, ?,
                    'created',
                    ?, ?, ?, ?, ?
                )
                """,
                (
                    key,
                    service,
                    title,
                    customer_name,
                    customer_id,
                    amount,
                    currency,
                    payment_method,
                    document["version"],
                    document["filename"],
                    document["page_count"],
                    current_time,
                    current_time,
                ),
            )

        conn.commit()

    return (
        get_payment(
            service,
            title,
        )
        or {}
    )


def update_payment_record(
    service: str,
    title: str,
    status: str,
    notes: str = "",
) -> Dict[str, Any]:

    key = business_key(
        service,
        title,
    )

    current_time = now_iso()

    with db_connect(PAYMENT_DB_PATH) as conn:

        if status == "reported":

            conn.execute(
                """
                UPDATE payment_orders

                SET payment_status = ?,

                    notes = ?,

                    reported_at = ?,

                    updated_at = ?

                WHERE business_key = ?
                """,
                (
                    status,
                    notes,
                    current_time,
                    current_time,
                    key,
                ),
            )

        elif status == "verified":

            conn.execute(
                """
                UPDATE payment_orders

                SET payment_status = ?,

                    notes = ?,

                    verified_at = ?,

                    updated_at = ?

                WHERE business_key = ?
                """,
                (
                    status,
                    notes,
                    current_time,
                    current_time,
                    key,
                ),
            )

        else:

            conn.execute(
                """
                UPDATE payment_orders

                SET payment_status = ?,

                    notes = ?,

                    updated_at = ?

                WHERE business_key = ?
                """,
                (
                    status,
                    notes,
                    current_time,
                    key,
                ),
            )

        conn.commit()

    return (
        get_payment(
            service,
            title,
        )
        or {}
    )


# ============================================================
# PUBLIC PRODUCT RESPONSE
# ============================================================

def product_public(
    product: Optional[Dict[str, Any]],
    payment: Optional[Dict[str, Any]],
) -> Dict[str, Any]:

    if not product:

        return {
            "found": False,
            "download_unlocked": False,
            "payment_verified": False,
        }

    payment_status = clean_text(
        (payment or {}).get(
            "payment_status"
        )
    ).casefold()

    return {

        "found": True,

        "service":
            product["service"],

        "document_title":
            product["document_title"],

        "customer_name":
            product.get("customer_name") or "",

        "customer_id":
            product.get("customer_id") or "",

        "amount":
            product.get("amount") or 0,

        "currency":
            product.get("currency") or CURRENCY,

        "document_version":
            product.get("document_version") or "",

        "document_filename":
            product.get("document_filename") or "",

        "filename":
            product.get("document_filename") or "",

        "pages":
            product.get("document_pages") or 1,

        "document_saved":
            bool(
                product.get(
                    "document_saved_path"
                )
            ),

        "document_saved_at":
            product.get(
                "document_saved_at"
            ),

        "download_unlocked":
            bool(
                product.get(
                    "download_unlocked"
                )
            ),

        "download_activated":
            bool(
                product.get(
                    "download_unlocked"
                )
            ),

        "download_count":
            product.get(
                "download_count"
            )
            or 0,

        "payment_status":
            payment_status
            or None,

        "payment_reported":
            payment_status == "reported",

        "payment_verified":
            payment_status == "verified",

        "activated_at":
            product.get(
                "activated_at"
            ),

        "created_at":
            product.get(
                "created_at"
            ),

        "updated_at":
            product.get(
                "updated_at"
            ),
    }


# ============================================================
# BACK OFFICE AUTH
# ============================================================

def check_back_office_key(
    x_back_office_key: Optional[str],
    query_key: Optional[str],
) -> bool:

    supplied = clean_text(
        x_back_office_key
        or query_key
    )

    return supplied == BACK_OFFICE_ADMIN_KEY


def require_back_office(
    x_back_office_key: Optional[str],
    query_key: Optional[str],
) -> None:

    if not check_back_office_key(
        x_back_office_key,
        query_key,
    ):

        raise HTTPException(
            status_code=401,
            detail="INVALID_BACK_OFFICE_KEY",
        )


# ============================================================
# DELIVERY URL HELPERS
# ============================================================

def api_base_url(
    request: Request,
) -> str:

    if PUBLIC_API_BASE_URL:
        return PUBLIC_API_BASE_URL

    return str(
        request.base_url
    ).rstrip("/")


def make_download_url(
    request: Request,
    service: str,
    title: str,
) -> str:

    query = urllib.parse.urlencode(
        {
            "service": service,
            "title": title,
        }
    )

    return (
        f"{api_base_url(request)}"
        f"/api/download?{query}"
    )


# ============================================================
# DELIVERY CHANNELS
# ============================================================

def make_delivery_channels(
    request: Request,
    product: Dict[str, Any],
) -> List[Dict[str, Any]]:

    service = clean_text(
        product["service"]
    )

    title = clean_text(
        product["document_title"]
    )

    download_url = make_download_url(
        request,
        service,
        title,
    )

    share_text = (
        "Naija Pocket Business Center document: "
        f"{title}\n"
        f"Service: {service}\n"
        "Download the exact saved approved document:"
    )

    whatsapp_text = (
        f"{share_text}\n"
        f"{download_url}"
    )

    email_subject = (
        "Naija Pocket Business Center — "
        f"{title}"
    )

    email_body = (
        f"{share_text}\n\n"
        f"{download_url}"
    )

    whatsapp_url = (
        "https://wa.me/?"
        + urllib.parse.urlencode(
            {
                "text":
                    whatsapp_text
            }
        )
    )

    email_url = (
        "mailto:?"
        + urllib.parse.urlencode(
            {
                "subject":
                    email_subject,

                "body":
                    email_body,
            }
        )
    )

    telegram_url = (
        "https://t.me/share/url?"
        + urllib.parse.urlencode(
            {
                "url":
                    download_url,

                "text":
                    share_text,
            }
        )
    )

    google_drive_url = (
        "https://drive.google.com/drive/my-drive"
    )

    unlocked = bool(
        product.get(
            "download_unlocked"
        )
    )

    return [

        {
            "id":
                "download",

            "channel":
                "phone",

            "name":
                "Download to Phone",

            "short_name":
                "Phone",

            "description":
                "Download the exact saved document directly to this device.",

            "action":
                "download",

            "action_url":
                download_url,

            "download_url":
                download_url,

            "available":
                unlocked,

            "requires_external_service":
                False,
        },

        {
            "id":
                "whatsapp",

            "channel":
                "whatsapp",

            "name":
                "WhatsApp",

            "short_name":
                "WhatsApp",

            "description":
                "Share the exact saved-document download link through WhatsApp.",

            "action":
                "share",

            "action_url":
                whatsapp_url,

            "download_url":
                download_url,

            "available":
                unlocked,

            "requires_external_service":
                True,
        },

        {
            "id":
                "email",

            "channel":
                "email",

            "name":
                "Email",

            "short_name":
                "Email",

            "description":
                "Open your email app with the exact saved-document link prepared.",

            "action":
                "share",

            "action_url":
                email_url,

            "download_url":
                download_url,

            "available":
                unlocked,

            "requires_external_service":
                True,
        },

        {
            "id":
                "telegram",

            "channel":
                "telegram",

            "name":
                "Telegram",

            "short_name":
                "Telegram",

            "description":
                "Share the exact saved-document link through Telegram.",

            "action":
                "share",

            "action_url":
                telegram_url,

            "download_url":
                download_url,

            "available":
                unlocked,

            "requires_external_service":
                True,
        },

        {
            "id":
                "google_drive",

            "channel":
                "google_drive",

            "name":
                "Google Drive",

            "short_name":
                "Google Drive",

            "description":
                "Open Google Drive, then upload the exact downloaded document.",

            "action":
                "open_drive",

            "action_url":
                google_drive_url,

            "download_url":
                download_url,

            "available":
                unlocked,

            "requires_external_service":
                True,

            "automatic_upload":
                False,
        },
    ]


def get_delivery_channel(
    request: Request,
    product: Dict[str, Any],
    channel: str,
) -> Optional[Dict[str, Any]]:

    channel = normalize_key_part(
        channel
    )

    aliases = {

        "phone":
            "download",

        "direct":
            "download",

        "download":
            "download",

        "whatsapp":
            "whatsapp",

        "email":
            "email",

        "telegram":
            "telegram",

        "google drive":
            "google_drive",

        "google_drive":
            "google_drive",

        "drive":
            "google_drive",
    }

    normalized = aliases.get(
        channel,
        channel,
    )

    for item in make_delivery_channels(
        request,
        product,
    ):

        if item["id"] == normalized:
            return item

    return None


# ============================================================
# DELIVERY LOGGING
# ============================================================

def log_delivery_event(
    product: Dict[str, Any],
    channel: str,
    recipient: str,
    action_url: str,
    status: str,
    message: str,
) -> None:

    with db_connect(
        PRODUCT_DB_PATH
    ) as conn:

        conn.execute(
            """
            INSERT INTO delivery_events (

                business_key,

                service,

                document_title,

                channel,

                recipient,

                action_url,

                status,

                message,

                created_at

            )

            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                business_key(
                    product["service"],
                    product["document_title"],
                ),

                product["service"],

                product["document_title"],

                channel,

                recipient,

                action_url,

                status,

                message,

                now_iso(),
            ),
        )

        conn.commit()


# ============================================================
# REQUEST MODELS
# ============================================================

class PaymentCreateRequest(BaseModel):

    customer_name: str = ""

    customer_id: str = ""

    service: str

    document_title: str

    amount: float = 0

    currency: str = "NGN"

    payment_method: str = "bank_transfer"

    document_version: str = ""

    version_id: str = ""

    pages: Any = None

    document_pages: Any = None

    document_text: str = ""

    filename: str = ""

    document_filename: str = ""

    notes: str = ""


class PaymentReportRequest(BaseModel):

    service: str

    document_title: str

    note: str = ""


class PaymentCompleteRequest(BaseModel):

    service: str

    document_title: str

    note: str = ""


class BackOfficeActivateRequest(BaseModel):

    service: str

    document_title: str

    verified: bool = True


class BackOfficeRejectRequest(BaseModel):

    service: str

    document_title: str

    reason: str = ""


class DeliveryRequest(BaseModel):

    service: str

    document_title: str

    channel: str

    recipient: str = ""

    note: str = ""


# ============================================================
# STARTUP
# ============================================================

@app.on_event("startup")
def startup_event() -> None:
    init_databases()


# ============================================================
# CREATE PAYMENT
# ============================================================

@app.post("/api/payment/create")
def payment_create(
    payload: PaymentCreateRequest,
) -> Dict[str, Any]:

    service = clean_text(
        payload.service
    )

    title = clean_text(
        payload.document_title
    )

    if not service:

        raise HTTPException(
            status_code=400,
            detail="SERVICE_REQUIRED",
        )

    if not title:

        raise HTTPException(
            status_code=400,
            detail="DOCUMENT_TITLE_REQUIRED",
        )

    if payload.amount < 0:

        raise HTTPException(
            status_code=400,
            detail="INVALID_AMOUNT",
        )

    raw = payload.model_dump()

    document = normalize_document_payload(
        raw
    )

    if not document["text"]:

        raise HTTPException(
            status_code=400,
            detail="DOCUMENT_CONTENT_REQUIRED",
        )

    product = upsert_product(

        service=service,

        title=title,

        customer_name=
            clean_text(
                payload.customer_name
            ),

        customer_id=
            clean_text(
                payload.customer_id
            ),

        amount=
            float(
                payload.amount
            ),

        currency=
            clean_text(
                payload.currency
            )
            or CURRENCY,

        document=
            document,

        notes=
            clean_text(
                payload.notes
            ),
    )

    # Save the exact reviewed document snapshot.
    try:

        saved_path = (
            save_exact_document_snapshot(
                service,
                title,
                document,
            )
        )

    except Exception as exc:

        raise HTTPException(
            status_code=500,
            detail=(
                "DOCUMENT_SAVE_FAILED: "
                f"{exc}"
            ),
        )

    with db_connect(
        PRODUCT_DB_PATH
    ) as conn:

        conn.execute(
            """
            UPDATE document_products

            SET document_saved_path = ?,

                document_saved_at =
                    COALESCE(
                        document_saved_at,
                        ?
                    ),

                updated_at = ?

            WHERE business_key = ?
            """,
            (
                saved_path,
                now_iso(),
                now_iso(),
                business_key(
                    service,
                    title,
                ),
            ),
        )

        conn.commit()

    # Payment record is separate from the document.
    try:

        payment = ensure_payment_record(

            service=service,

            title=title,

            customer_name=
                clean_text(
                    payload.customer_name
                ),

            customer_id=
                clean_text(
                    payload.customer_id
                ),

            amount=
                float(
                    payload.amount
                ),

            currency=
                clean_text(
                    payload.currency
                )
                or CURRENCY,

            payment_method=
                clean_text(
                    payload.payment_method
                )
                or "bank_transfer",

            document=
                document,
        )

    except Exception:

        # A payment DB issue must never
        # destroy the saved document.
        payment = (
            get_payment(
                service,
                title,
            )
            or {}
        )

    product = (
        get_product(
            service,
            title,
        )
        or product
    )

    return {

        "ok":
            True,

        "message":
            "Exact reviewed document saved and payment prepared.",

        "product":
            product_public(
                product,
                payment,
            ),

        "saved_document":
            {
                "saved":
                    True,

                "filename":
                    product.get(
                        "document_filename"
                    ),

                "pages":
                    product.get(
                        "document_pages"
                    )
                    or 1,

                "path_saved":
                    bool(
                        product.get(
                            "document_saved_path"
                        )
                    ),
            },
    }


# ============================================================
# REPORT PAYMENT
# ============================================================

@app.post("/api/payment/report")
def payment_report(
    payload: PaymentReportRequest,
) -> Dict[str, Any]:

    service = clean_text(
        payload.service
    )

    title = clean_text(
        payload.document_title
    )

    product = get_product_or_single(
        service,
        title,
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    actual_service = product[
        "service"
    ]

    actual_title = product[
        "document_title"
    ]

    saved_path = clean_text(
        product.get(
            "document_saved_path"
        )
    )

    if (
        not saved_path
        or not Path(saved_path).exists()
    ):

        document = normalize_document_payload(
            {
                "document_text":
                    product.get(
                        "document_payload"
                    )
                    or "",

                "filename":
                    product.get(
                        "document_filename"
                    )
                    or "",

                "pages":
                    product.get(
                        "document_payload"
                    )
                    or "",

                "document_version":
                    product.get(
                        "document_version"
                    )
                    or "",
            }
        )

        if not document["text"]:

            raise HTTPException(
                status_code=409,
                detail=
                    "SAVED_DOCUMENT_CONTENT_MISSING",
            )

        saved_path = (
            save_exact_document_snapshot(
                actual_service,
                actual_title,
                document,
            )
        )

    payment = get_payment(
        actual_service,
        actual_title,
    )

    if not payment:

        payment = ensure_payment_record(

            service=
                actual_service,

            title=
                actual_title,

            customer_name=
                product.get(
                    "customer_name"
                )
                or "",

            customer_id=
                product.get(
                    "customer_id"
                )
                or "",

            amount=
                float(
                    product.get(
                        "amount"
                    )
                    or 0
                ),

            currency=
                product.get(
                    "currency"
                )
                or CURRENCY,

            payment_method=
                "bank_transfer",

            document=
                {
                    "version":
                        product.get(
                            "document_version"
                        )
                        or "",

                    "filename":
                        product.get(
                            "document_filename"
                        )
                        or "",

                    "page_count":
                        product.get(
                            "document_pages"
                        )
                        or 1,

                    "text":
                        product.get(
                            "document_payload"
                        )
                        or "",

                    "pages":
                        [
                            product.get(
                                "document_payload"
                            )
                            or ""
                        ],
                },
        )

    payment = update_payment_record(

        actual_service,

        actual_title,

        "reported",

        clean_text(
            payload.note
        ),
    )

    product = (
        get_product(
            actual_service,
            actual_title,
        )
        or product
    )

    return {

        "ok":
            True,

        "message":
            "Payment reported. Customer Care will verify it.",

        "product":
            product_public(
                product,
                payment,
            ),
    }


# ============================================================
# PAYMENT STATUS
# ============================================================

@app.get("/api/payment/status")
def payment_status(
    service: str = Query(...),
    title: str = Query(...),
) -> Dict[str, Any]:

    product = get_product_or_single(
        service,
        title,
    )

    if not product:

        return {

            "ok":
                True,

            "found":
                False,

            "download_unlocked":
                False,
        }

    payment = get_payment(
        product["service"],
        product["document_title"],
    )

    return {

        "ok":
            True,

        "product":
            product_public(
                product,
                payment,
            ),
    }


# ============================================================
# PAYMENT COMPLETE
# ============================================================

@app.post("/api/payment/complete")
def payment_complete(
    payload: PaymentCompleteRequest,
) -> Dict[str, Any]:

    product = get_product_or_single(

        clean_text(
            payload.service
        ),

        clean_text(
            payload.document_title
        ),
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    payment = update_payment_record(

        product["service"],

        product["document_title"],

        "verified",

        clean_text(
            payload.note
        ),
    )

    product = (
        get_product(
            product["service"],
            product["document_title"],
        )
        or product
    )

    return {

        "ok":
            True,

        "message":
            "Payment marked verified. Download activation remains a separate Customer Care action.",

        "product":
            product_public(
                product,
                payment,
            ),
    }


# ============================================================
# CUSTOMER CARE PAYMENTS
# ============================================================

@app.get("/api/customer-care/payments")
def customer_care_payments(

    x_back_office_key:
        Optional[str] =
            Header(default=None),

    key:
        Optional[str] =
            Query(default=None),

) -> Dict[str, Any]:

    require_back_office(
        x_back_office_key,
        key,
    )

    with db_connect(
        PRODUCT_DB_PATH
    ) as conn:

        products = [

            row_dict(row)

            for row in conn.execute(
                """
                SELECT *
                FROM document_products

                ORDER BY
                    updated_at DESC,
                    id DESC
                """
            ).fetchall()
        ]

    records = []

    for product in products:

        payment = get_payment(

            product["service"],

            product["document_title"],
        )

        records.append(
            product_public(
                product,
                payment,
            )
        )

    return {

        "ok":
            True,

        "payments":
            records,

        "records":
            records,

        "items":
            records,
    }


# ============================================================
# CUSTOMER CARE VERIFY
# ============================================================

@app.post(
    "/api/customer-care/payment/verify"
)
def customer_care_payment_verify(

    payload:
        BackOfficeActivateRequest,

    x_back_office_key:
        Optional[str] =
            Header(default=None),

    key:
        Optional[str] =
            Query(default=None),

) -> Dict[str, Any]:

    require_back_office(
        x_back_office_key,
        key,
    )

    product = get_product_or_single(

        payload.service,

        payload.document_title,
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    payment = update_payment_record(

        product["service"],

        product["document_title"],

        "verified",

        "Verified by Customer Care.",
    )

    product = (
        get_product(
            product["service"],
            product["document_title"],
        )
        or product
    )

    return {

        "ok":
            True,

        "message":
            "Payment verified. Download still requires explicit activation.",

        "product":
            product_public(
                product,
                payment,
            ),
    }


# ============================================================
# BACK OFFICE LOGIN
# ============================================================

@app.post("/api/back-office/login")
def back_office_login(

    x_back_office_key:
        Optional[str] =
            Header(default=None),

    key:
        Optional[str] =
            Query(default=None),

) -> Dict[str, Any]:

    require_back_office(
        x_back_office_key,
        key,
    )

    return {

        "ok":
            True,

        "authenticated":
            True,

        "message":
            "Customer Care Back Office access granted.",
    }


# ============================================================
# BACK OFFICE SAVED DOCUMENTS
# ============================================================

@app.get("/api/back-office/payments")
def back_office_payments(

    x_back_office_key:
        Optional[str] =
            Header(default=None),

    key:
        Optional[str] =
            Query(default=None),

) -> Dict[str, Any]:

    require_back_office(
        x_back_office_key,
        key,
    )

    with db_connect(
        PRODUCT_DB_PATH
    ) as conn:

        products = [

            row_dict(row)

            for row in conn.execute(
                """
                SELECT *
                FROM document_products

                ORDER BY
                    updated_at DESC,
                    id DESC
                """
            ).fetchall()
        ]

    records = []

    for product in products:

        payment = get_payment(

            product["service"],

            product["document_title"],
        )

        record = product_public(
            product,
            payment,
        )

        record["saved"] = True

        record["saved_document_path"] = (
            product.get(
                "document_saved_path"
            )
        )

        record["document_saved_path"] = (
            product.get(
                "document_saved_path"
            )
        )

        record["verified"] = bool(

            payment

            and payment.get(
                "payment_status"
            ) == "verified"
        )

        record["reported"] = bool(

            payment

            and payment.get(
                "payment_status"
            ) == "reported"
        )

        record["activated"] = bool(
            product.get(
                "download_unlocked"
            )
        )

        records.append(
            record
        )

    return {

        "ok":
            True,

        "payments":
            records,

        "records":
            records,

        "items":
            records,
    }


# ============================================================
# BACK OFFICE JOBS
# ============================================================

@app.get("/api/back-office/jobs")
def back_office_jobs(

    x_back_office_key:
        Optional[str] =
            Header(default=None),

    key:
        Optional[str] =
            Query(default=None),

) -> Dict[str, Any]:

    result = back_office_payments(
        x_back_office_key=
            x_back_office_key,

        key=
            key,
    )

    return {

        "ok":
            True,

        "jobs":
            result["payments"],

        "items":
            result["payments"],
    }


# ============================================================
# BACK OFFICE SINGLE PAYMENT
# ============================================================

@app.get("/api/back-office/payment")
def back_office_payment(

    service: str,

    title: str,

    x_back_office_key:
        Optional[str] =
            Header(default=None),

    key:
        Optional[str] =
            Query(default=None),

) -> Dict[str, Any]:

    require_back_office(
        x_back_office_key,
        key,
    )

    product = get_product_or_single(
        service,
        title,
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    payment = get_payment(

        product["service"],

        product["document_title"],
    )

    return {

        "ok":
            True,

        "product":
            product_public(
                product,
                payment,
            ),

        "payment":
            payment,
    }


# ============================================================
# BACK OFFICE DOCUMENT INFO
# ============================================================

@app.get(
    "/api/back-office/document-info"
)
def back_office_document_info(

    service: str,

    title: str,

    x_back_office_key:
        Optional[str] =
            Header(default=None),

    key:
        Optional[str] =
            Query(default=None),

) -> Dict[str, Any]:

    require_back_office(
        x_back_office_key,
        key,
    )

    product = get_product_or_single(
        service,
        title,
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    saved_path = clean_text(
        product.get(
            "document_saved_path"
        )
    )

    return {

        "ok":
            True,

        "service":
            product["service"],

        "document_title":
            product["document_title"],

        "filename":
            product.get(
                "document_filename"
            )
            or "",

        "pages":
            product.get(
                "document_pages"
            )
            or 1,

        "document_version":
            product.get(
                "document_version"
            )
            or "",

        "saved":
            bool(
                saved_path
                and Path(saved_path).exists()
            ),

        "saved_path":
            saved_path,

        "saved_at":
            product.get(
                "document_saved_at"
            ),

        "download_unlocked":
            bool(
                product.get(
                    "download_unlocked"
                )
            ),
    }


# ============================================================
# BACK OFFICE DOCUMENT CONTENT
# ============================================================

@app.get("/api/back-office/document")
def back_office_document(

    service: str,

    title: str,

    x_back_office_key:
        Optional[str] =
            Header(default=None),

    key:
        Optional[str] =
            Query(default=None),

) -> Dict[str, Any]:

    require_back_office(
        x_back_office_key,
        key,
    )

    product = get_product_or_single(
        service,
        title,
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    saved_path = clean_text(
        product.get(
            "document_saved_path"
        )
    )

    return {

        "ok":
            True,

        "service":
            product["service"],

        "document_title":
            product["document_title"],

        "filename":
            product.get(
                "document_filename"
            )
            or "",

        "pages":
            product.get(
                "document_pages"
            )
            or 1,

        "document_version":
            product.get(
                "document_version"
            )
            or "",

        "document_text":
            product.get(
                "document_payload"
            )
            or "",

        "saved":
            bool(
                saved_path
                and Path(saved_path).exists()
            ),

        "saved_path":
            saved_path,

        "download_unlocked":
            bool(
                product.get(
                    "download_unlocked"
                )
            ),
    }


# ============================================================
# BACK OFFICE PAYMENT VERIFY
# ============================================================

@app.post(
    "/api/back-office/payment/verify"
)
def back_office_payment_verify(

    payload:
        BackOfficeActivateRequest,

    x_back_office_key:
        Optional[str] =
            Header(default=None),

    key:
        Optional[str] =
            Query(default=None),

) -> Dict[str, Any]:

    require_back_office(
        x_back_office_key,
        key,
    )

    product = get_product_or_single(

        payload.service,

        payload.document_title,
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    payment = update_payment_record(

        product["service"],

        product["document_title"],

        "verified",

        "Payment verified by Customer Care.",
    )

    product = (
        get_product(
            product["service"],
            product["document_title"],
        )
        or product
    )

    return {

        "ok":
            True,

        "message":
            "Payment verified. Download activation remains separate.",

        "product":
            product_public(
                product,
                payment,
            ),
    }


# ============================================================
# ACTIVATE DOWNLOAD
# ============================================================

@app.post(
    "/api/back-office/activate-download"
)
def back_office_activate_download(

    payload:
        BackOfficeActivateRequest,

    x_back_office_key:
        Optional[str] =
            Header(default=None),

    key:
        Optional[str] =
            Query(default=None),

) -> Dict[str, Any]:

    require_back_office(
        x_back_office_key,
        key,
    )

    if not payload.verified:

        raise HTTPException(
            status_code=400,
            detail=
                "PAYMENT_MUST_BE_VERIFIED_BEFORE_ACTIVATION",
        )

    product = get_product_or_single(

        payload.service,

        payload.document_title,
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    payment = get_payment(

        product["service"],

        product["document_title"],
    )

    if (
        not payment
        or clean_text(
            payment.get(
                "payment_status"
            )
        ).casefold()
        != "verified"
    ):

        raise HTTPException(
            status_code=409,
            detail="PAYMENT_NOT_VERIFIED",
        )

    saved_path = clean_text(
        product.get(
            "document_saved_path"
        )
    )

    if (
        not saved_path
        or not Path(saved_path).exists()
    ):

        raise HTTPException(
            status_code=409,
            detail="SAVED_DOCUMENT_FILE_MISSING",
        )

    activated = activate_product_download(

        product["service"],

        product["document_title"],
    )

    payment = get_payment(

        activated["service"],

        activated["document_title"],
    )

    return {

        "ok":
            True,

        "message":
            "Download activated. Customer can now receive the exact saved document.",

        "product":
            product_public(
                activated,
                payment,
            ),
    }


# ============================================================
# BACK OFFICE REJECT
# ============================================================

@app.post(
    "/api/back-office/payment/reject"
)
def back_office_payment_reject(

    payload:
        BackOfficeRejectRequest,

    x_back_office_key:
        Optional[str] =
            Header(default=None),

    key:
        Optional[str] =
            Query(default=None),

) -> Dict[str, Any]:

    require_back_office(
        x_back_office_key,
        key,
    )

    product = get_product_or_single(

        payload.service,

        payload.document_title,
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    payment = update_payment_record(

        product["service"],

        product["document_title"],

        "rejected",

        clean_text(
            payload.reason
        ),
    )

    product = (
        get_product(
            product["service"],
            product["document_title"],
        )
        or product
    )

    return {

        "ok":
            True,

        "message":
            "Payment marked rejected. Download remains locked.",

        "product":
            product_public(
                product,
                payment,
            ),
    }


# ============================================================
# CUSTOMER DOWNLOAD
# ============================================================

@app.get("/api/download")
def download_document(

    service: str,

    title: str,

) -> FileResponse:

    product = get_product_or_single(
        service,
        title,
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    if not bool(
        product.get(
            "download_unlocked"
        )
    ):

        raise HTTPException(
            status_code=403,
            detail="DOWNLOAD_NOT_ACTIVATED",
        )

    saved_path = clean_text(
        product.get(
            "document_saved_path"
        )
    )

    if not saved_path:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_FILE_MISSING",
        )

    path = Path(
        saved_path
    )

    if (
        not path.exists()
        or not path.is_file()
    ):

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_FILE_MISSING",
        )

    current_time = now_iso()

    with db_connect(
        PRODUCT_DB_PATH
    ) as conn:

        conn.execute(
            """
            UPDATE document_products

            SET download_count =
                    COALESCE(
                        download_count,
                        0
                    ) + 1,

                downloaded_at = ?,

                updated_at = ?

            WHERE business_key = ?
            """,
            (
                current_time,
                current_time,
                business_key(
                    product["service"],
                    product["document_title"],
                ),
            ),
        )

        conn.commit()

    suffix = path.suffix.casefold()

    if suffix == ".docx":

        media_type = (
            "application/"
            "vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )

    elif suffix == ".pdf":

        media_type = "application/pdf"

    else:

        media_type = (
            "application/octet-stream"
        )

    return FileResponse(

        path=str(path),

        filename=path.name,

        media_type=media_type,
    )


# ============================================================
# CUSTOMER DELIVERY CHANNELS
# ============================================================

@app.get("/api/delivery/channels")
def delivery_channels(

    request: Request,

    service: str,

    title: str,

) -> Dict[str, Any]:

    product = get_product_or_single(
        service,
        title,
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    saved_path = clean_text(
        product.get(
            "document_saved_path"
        )
    )

    if (
        not saved_path
        or not Path(saved_path).exists()
    ):

        raise HTTPException(
            status_code=409,
            detail="SAVED_DOCUMENT_FILE_MISSING",
        )

    channels = make_delivery_channels(
        request,
        product,
    )

    return {

        "ok":
            True,

        "service":
            product["service"],

        "document_title":
            product["document_title"],

        "filename":
            product.get(
                "document_filename"
            )
            or Path(
                saved_path
            ).name,

        "download_unlocked":
            bool(
                product.get(
                    "download_unlocked"
                )
            ),

        "download_url":
            make_download_url(

                request,

                product["service"],

                product["document_title"],
            ),

        "channels":
            channels,
    }


# ============================================================
# CUSTOMER DELIVERY PREPARE
# ============================================================

@app.post("/api/delivery/prepare")
def delivery_prepare(

    request: Request,

    payload: DeliveryRequest,

) -> Dict[str, Any]:

    product = get_product_or_single(

        payload.service,

        payload.document_title,
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    if not bool(
        product.get(
            "download_unlocked"
        )
    ):

        raise HTTPException(
            status_code=403,
            detail="DOWNLOAD_NOT_ACTIVATED",
        )

    saved_path = clean_text(
        product.get(
            "document_saved_path"
        )
    )

    if (
        not saved_path
        or not Path(saved_path).exists()
    ):

        raise HTTPException(
            status_code=409,
            detail="SAVED_DOCUMENT_FILE_MISSING",
        )

    channel = get_delivery_channel(

        request,

        product,

        payload.channel,
    )

    if not channel:

        raise HTTPException(
            status_code=400,
            detail=
                "DELIVERY_CHANNEL_NOT_SUPPORTED",
        )

    log_delivery_event(

        product=product,

        channel=channel["id"],

        recipient=
            clean_text(
                payload.recipient
            ),

        action_url=
            channel["action_url"],

        status="prepared",

        message=
            clean_text(
                payload.note
            )
            or channel["description"],
    )

    return {

        "ok":
            True,

        "message":
            "Delivery action prepared.",

        "service":
            product["service"],

        "document_title":
            product["document_title"],

        "channel":
            channel,
    }


# ============================================================
# OPTIONAL LEGACY WEBHOOKS
# ============================================================

def webhook_for_channel(
    channel: str,
) -> str:

    mapping = {

        "email":
            DELIVERY_EMAIL_WEBHOOK,

        "whatsapp":
            DELIVERY_WHATSAPP_WEBHOOK,

        "telegram":
            DELIVERY_TELEGRAM_WEBHOOK,

        "google_drive":
            DELIVERY_GOOGLE_DRIVE_WEBHOOK,
    }

    return mapping.get(
        normalize_key_part(
            channel
        ),
        "",
    )


def send_delivery_webhook(
    webhook_url: str,
    payload: Dict[str, Any],
) -> Dict[str, Any]:

    if not webhook_url:

        return {

            "sent":
                False,

            "configured":
                False,

            "message":
                "No delivery webhook configured.",
        }

    try:

        import json
        import urllib.request

        request_data = json.dumps(
            payload
        ).encode("utf-8")

        request = (
            urllib.request.Request(

                webhook_url,

                data=request_data,

                headers={
                    "Content-Type":
                        "application/json"
                },

                method="POST",
            )
        )

        with urllib.request.urlopen(
            request,
            timeout=15,
        ) as response:

            body = response.read().decode(
                "utf-8",
                errors="replace",
            )

            status_code = getattr(
                response,
                "status",
                200,
            )

        return {

            "sent":
                True,

            "configured":
                True,

            "status_code":
                status_code,

            "response":
                body[:2000],
        }

    except Exception as exc:

        return {

            "sent":
                False,

            "configured":
                True,

            "message":
                str(exc),
        }


# ============================================================
# BACK OFFICE DELIVERY
# ============================================================

@app.post("/api/back-office/delivery")
def back_office_delivery(

    request: Request,

    payload: DeliveryRequest,

    x_back_office_key:
        Optional[str] =
            Header(default=None),

    key:
        Optional[str] =
            Query(default=None),

) -> Dict[str, Any]:

    require_back_office(
        x_back_office_key,
        key,
    )

    product = get_product_or_single(

        payload.service,

        payload.document_title,
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    channel = normalize_key_part(
        payload.channel
    )

    if channel in {
        "direct",
        "download",
        "phone",
    }:

        if not bool(
            product.get(
                "download_unlocked"
            )
        ):

            raise HTTPException(
                status_code=403,
                detail="DOWNLOAD_NOT_ACTIVATED",
            )

        item = get_delivery_channel(

            request,

            product,

            "download",
        )

        log_delivery_event(

            product,

            "download",

            clean_text(
                payload.recipient
            ),

            item["action_url"]
            if item
            else "",

            "prepared",

            clean_text(
                payload.note
            )
            or "Direct download prepared.",
        )

        return {

            "ok":
                True,

            "message":
                "Direct download prepared.",

            "channel":
                item,
        }

    if not bool(
        product.get(
            "download_unlocked"
        )
    ):

        raise HTTPException(
            status_code=403,
            detail="DOWNLOAD_NOT_ACTIVATED",
        )

    item = get_delivery_channel(

        request,

        product,

        channel,
    )

    if not item:

        raise HTTPException(
            status_code=400,
            detail=
                "DELIVERY_CHANNEL_NOT_SUPPORTED",
        )

    webhook_url = webhook_for_channel(
        channel
    )

    # Existing webhook functionality is preserved.
    if webhook_url:

        result = send_delivery_webhook(

            webhook_url,

            {
                "service":
                    product["service"],

                "document_title":
                    product["document_title"],

                "filename":
                    product.get(
                        "document_filename"
                    )
                    or "",

                "saved_document_path":
                    product.get(
                        "document_saved_path"
                    )
                    or "",

                "download_url":
                    item["download_url"],

                "recipient":
                    clean_text(
                        payload.recipient
                    ),

                "note":
                    clean_text(
                        payload.note
                    ),

                "channel":
                    channel,
            },
        )

        status = (
            "webhook_sent"
            if result.get("sent")
            else "webhook_failed"
        )

        log_delivery_event(

            product,

            channel,

            clean_text(
                payload.recipient
            ),

            item["action_url"],

            status,

            result.get(
                "message",
                "",
            ),
        )

        return {

            "ok":
                bool(
                    result.get(
                        "sent"
                    )
                ),

            "message":
                (
                    "Delivery webhook sent."
                    if result.get(
                        "sent"
                    )
                    else
                    "Delivery webhook could not be sent."
                ),

            "channel":
                item,

            "webhook":
                result,
        }

    # Standard share action works even
    # without external API credentials.
    log_delivery_event(

        product,

        channel,

        clean_text(
            payload.recipient
        ),

        item["action_url"],

        "prepared",

        "Standard share action prepared.",
    )

    return {

        "ok":
            True,

        "message":
            "Delivery action prepared.",

        "channel":
            item,

        "webhook":
            {
                "sent":
                    False,

                "configured":
                    False,

                "message":
                    "Standard share action is available.",
            },
    }


# ============================================================
# BACK OFFICE DELIVERY FILE
# ============================================================

@app.get(
    "/api/back-office/delivery-file"
)
def back_office_delivery_file(

    service: str,

    title: str,

    x_back_office_key:
        Optional[str] =
            Header(default=None),

    key:
        Optional[str] =
            Query(default=None),

) -> FileResponse:

    require_back_office(
        x_back_office_key,
        key,
    )

    product = get_product_or_single(
        service,
        title,
    )

    if not product:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_NOT_FOUND",
        )

    saved_path = clean_text(
        product.get(
            "document_saved_path"
        )
    )

    if not saved_path:

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_FILE_MISSING",
        )

    path = Path(
        saved_path
    )

    if (
        not path.exists()
        or not path.is_file()
    ):

        raise HTTPException(
            status_code=404,
            detail="SAVED_DOCUMENT_FILE_MISSING",
        )

    suffix = path.suffix.casefold()

    if suffix == ".docx":

        media_type = (
            "application/"
            "vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )

    elif suffix == ".pdf":

        media_type = "application/pdf"

    else:

        media_type = (
            "application/octet-stream"
        )

    return FileResponse(

        path=str(path),

        filename=path.name,

        media_type=media_type,
    )


# ============================================================
# DELIVERY HISTORY
# ============================================================

@app.get(
    "/api/back-office/delivery-history"
)
def back_office_delivery_history(

    service: Optional[str] = None,

    title: Optional[str] = None,

    x_back_office_key:
        Optional[str] =
            Header(default=None),

    key:
        Optional[str] =
            Query(default=None),

) -> Dict[str, Any]:

    require_back_office(
        x_back_office_key,
        key,
    )

    with db_connect(
        PRODUCT_DB_PATH
    ) as conn:

        if service and title:

            rows = conn.execute(
                """
                SELECT *
                FROM delivery_events

                WHERE business_key = ?

                ORDER BY
                    created_at DESC,
                    id DESC
                """,
                (
                    business_key(
                        service,
                        title,
                    ),
                ),
            ).fetchall()

        else:

            rows = conn.execute(
                """
                SELECT *
                FROM delivery_events

                ORDER BY
                    created_at DESC,
                    id DESC

                LIMIT 200
                """
            ).fetchall()

    events = [
        row_dict(row)
        for row in rows
    ]

    return {

        "ok":
            True,

        "events":
            events,
    }


# ============================================================
# PRODUCT STATUS
# ============================================================

@app.get("/api/product/status")
def product_status(

    service: str,

    title: str,

) -> Dict[str, Any]:

    product = get_product_or_single(
        service,
        title,
    )

    if not product:

        return {

            "ok":
                True,

            "found":
                False,

            "download_unlocked":
                False,
        }

    payment = get_payment(

        product["service"],

        product["document_title"],
    )

    return {

        "ok":
            True,

        "product":
            product_public(
                product,
                payment,
            ),
    }


# ============================================================
# HEALTH
# ============================================================

@app.get("/health")
def health() -> Dict[str, Any]:

    init_databases()

    return {

        "ok":
            True,

        "service":
            "Naija Pocket Business Center Payment API",

        "version":
            APP_VERSION,

        "back_office_key_configured":
            True,

        "download_directory":
            str(
                DOWNLOAD_DIR
            ),

        "delivery_channels":
            [
                "phone",
                "whatsapp",
                "email",
                "telegram",
                "google_drive",
            ],
    }


# ============================================================
# ROOT
# ============================================================

@app.get("/")
def root() -> Dict[str, Any]:

    return {

        "ok":
            True,

        "service":
            "Naija Pocket Business Center Payment API",

        "version":
            APP_VERSION,

        "message":
            "Payment, saved-document and delivery service is running.",
    }


# ============================================================
# LOCAL RUN
# ============================================================

if __name__ == "__main__":

    import uvicorn

    init_databases()

    uvicorn.run(

        app,

        host="0.0.0.0",

        port=int(
            os.getenv(
                "PORT",
                "8000",
            )
        ),
    )
